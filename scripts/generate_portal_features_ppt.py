from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "FloodGuard_Portal_Features_and_Future_Scope.pptx"

WIDE = Inches(13.333)
HIGH = Inches(7.5)

NAVY = RGBColor(9, 31, 55)
INK = RGBColor(22, 33, 45)
MUTED = RGBColor(88, 102, 118)
TEAL = RGBColor(0, 148, 136)
BLUE = RGBColor(36, 114, 200)
ORANGE = RGBColor(230, 128, 52)
RED = RGBColor(202, 57, 57)
GREEN = RGBColor(67, 150, 102)
PALE = RGBColor(244, 248, 251)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(214, 224, 234)


def set_font(run, size=18, bold=False, color=INK):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, x, y, w, h, text="", size=18, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size=size, bold=bold, color=color)
    return box


def add_rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=False):
    shape_type = 5 if radius else 1
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_footer(slide, number):
    add_textbox(slide, Inches(0.55), Inches(7.12), Inches(4.4), Inches(0.22),
                "FloodGuard AI Disaster Management Portal", size=7, color=MUTED)
    add_textbox(slide, Inches(12.15), Inches(7.12), Inches(0.6), Inches(0.22),
                f"{number:02}", size=7, color=MUTED, align=PP_ALIGN.RIGHT)


def add_header(slide, kicker, title, number):
    add_textbox(slide, Inches(0.55), Inches(0.28), Inches(3.0), Inches(0.24),
                kicker.upper(), size=8, bold=True, color=TEAL)
    add_textbox(slide, Inches(0.55), Inches(0.58), Inches(8.8), Inches(0.62),
                title, size=24, bold=True, color=NAVY)
    add_footer(slide, number)


def add_bullets(slide, items, x, y, w, h, size=13, color=INK, gap=0.06):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"- {item}"
        p.level = 0
        p.space_after = Pt(gap * 72)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def add_plain_list(slide, items, x, y, w, h, size=13, color=INK):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(5)
    return box


def add_card(slide, x, y, w, h, title, body, accent=TEAL):
    add_rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True)
    band = slide.shapes.add_shape(1, x, y, Inches(0.08), h)
    band.fill.solid()
    band.fill.fore_color.rgb = accent
    band.line.color.rgb = accent
    add_textbox(slide, x + Inches(0.18), y + Inches(0.15), w - Inches(0.35), Inches(0.28),
                title, size=11, bold=True, color=NAVY)
    add_textbox(slide, x + Inches(0.18), y + Inches(0.48), w - Inches(0.35), h - Inches(0.58),
                body, size=9.5, color=MUTED)


def add_metric(slide, x, y, w, label, value, color=TEAL):
    add_rect(slide, x, y, w, Inches(0.82), fill=WHITE, line=LINE, radius=True)
    add_textbox(slide, x + Inches(0.14), y + Inches(0.12), w - Inches(0.28), Inches(0.2),
                label.upper(), size=6.8, bold=True, color=MUTED)
    add_textbox(slide, x + Inches(0.14), y + Inches(0.33), w - Inches(0.28), Inches(0.32),
                value, size=17, bold=True, color=color)


def add_image_contain(slide, path, x, y, w, h):
    path = ROOT / path
    if not path.exists():
        add_rect(slide, x, y, w, h, fill=PALE, line=LINE, radius=True)
        add_textbox(slide, x + Inches(0.2), y + Inches(0.2), w - Inches(0.4), Inches(0.3),
                    f"Image missing: {path.name}", size=10, color=MUTED)
        return None
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pic_w = int(iw * scale)
    pic_h = int(ih * scale)
    px = x + int((w - pic_w) / 2)
    py = y + int((h - pic_h) / 2)
    add_rect(slide, x, y, w, h, fill=RGBColor(233, 239, 246), line=LINE, radius=True)
    pic = slide.shapes.add_picture(str(path), px, py, width=pic_w, height=pic_h)
    return pic


def add_image_cover(slide, path, x, y, w, h):
    path = ROOT / path
    with Image.open(path) as im:
        iw, ih = im.size
    box_ratio = w / h
    img_ratio = iw / ih
    pic = slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    if img_ratio > box_ratio:
        crop = 1 - box_ratio / img_ratio
        pic.crop_left = crop / 2
        pic.crop_right = crop / 2
    else:
        crop = 1 - img_ratio / box_ratio
        pic.crop_top = crop / 2
        pic.crop_bottom = crop / 2
    return pic


def add_timeline(slide, phases, x, y, w):
    segment_w = w / len(phases)
    for i, (label, title, body, color) in enumerate(phases):
        sx = x + int(segment_w * i)
        add_rect(slide, sx, y, int(segment_w) - Inches(0.1), Inches(3.4), fill=WHITE, line=LINE, radius=True)
        dot = slide.shapes.add_shape(9, sx + Inches(0.22), y + Inches(0.22), Inches(0.34), Inches(0.34))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.color.rgb = color
        add_textbox(slide, sx + Inches(0.68), y + Inches(0.17), int(segment_w) - Inches(0.9), Inches(0.22),
                    label, size=8, bold=True, color=color)
        add_textbox(slide, sx + Inches(0.22), y + Inches(0.7), int(segment_w) - Inches(0.45), Inches(0.35),
                    title, size=13, bold=True, color=NAVY)
        add_plain_list(slide, body, sx + Inches(0.24), y + Inches(1.14), int(segment_w) - Inches(0.48), Inches(2.0),
                       size=9.1, color=MUTED)


def build():
    prs = Presentation()
    prs.slide_width = WIDE
    prs.slide_height = HIGH
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_image_cover(slide, "frontend/river-monitoring-desktop.png", Inches(0), Inches(0), WIDE, HIGH)
    overlay = slide.shapes.add_shape(1, Inches(0), Inches(0), WIDE, HIGH)
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(5, 19, 34)
    overlay.fill.transparency = 18
    overlay.line.fill.background()
    add_textbox(slide, Inches(0.75), Inches(0.72), Inches(3.1), Inches(0.26),
                "PORTAL FEATURE DECK", size=9, bold=True, color=RGBColor(151, 236, 219))
    add_textbox(slide, Inches(0.75), Inches(1.18), Inches(7.8), Inches(1.3),
                "FloodGuard AI Disaster Management Portal", size=36, bold=True, color=WHITE)
    add_textbox(slide, Inches(0.78), Inches(2.65), Inches(6.5), Inches(0.62),
                "All current portal capabilities plus future scope and implementation roadmap",
                size=18, color=RGBColor(220, 235, 245))
    add_textbox(slide, Inches(0.78), Inches(6.82), Inches(5.6), Inches(0.28),
                "Prepared from the local disaster_management repository - 12 Oct 2026",
                size=8.5, color=RGBColor(204, 220, 232))

    # 2. Executive view
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Executive Summary", "What the Portal Delivers Today", 2)
    add_textbox(slide, Inches(0.72), Inches(1.38), Inches(4.9), Inches(0.52),
                "A command dashboard for flood intelligence, field reporting, simulation, and response coordination across Pakistan-focused operational geography.",
                size=17, bold=True, color=INK)
    add_metric(slide, Inches(0.72), Inches(2.25), Inches(2.15), "Main views", "10+", BLUE)
    add_metric(slide, Inches(3.02), Inches(2.25), Inches(2.15), "Priority sites", "28", TEAL)
    add_metric(slide, Inches(5.32), Inches(2.25), Inches(2.15), "Refresh modes", "15m / 3m", ORANGE)
    add_metric(slide, Inches(7.62), Inches(2.25), Inches(2.15), "ADM2 screen", "126 districts", GREEN)
    add_card(slide, Inches(0.72), Inches(3.36), Inches(3.9), Inches(1.42), "Operational awareness",
             "Overview map, KPIs, live weather and river context, hourly analytics, activity stream, alerts, and incidents.", BLUE)
    add_card(slide, Inches(4.86), Inches(3.36), Inches(3.9), Inches(1.42), "GIS intelligence",
             "District risk, rainfall raster, rivers, catchments, glaciers, Google Flood Hub points, reports, incidents, and resources.", TEAL)
    add_card(slide, Inches(9.0), Inches(3.36), Inches(3.4), Inches(1.42), "Decision support",
             "Predictive models, upstream forecasts, monitoring simulation, citizen reports, population exposure, and automation context.", ORANGE)
    add_bullets(slide, [
        "Production defaults avoid silent mock fallbacks; simulation is isolated and clearly labeled.",
        "Backend APIs expose dashboard summaries, GIS layers, forecasts, reports, incidents, resources, and analysis artifacts.",
        "Frontend is built as a responsive authority dashboard with MapLibre-based operational layers.",
    ], Inches(0.78), Inches(5.38), Inches(11.6), Inches(0.9), size=12)

    # 3. Navigation
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Portal Structure", "Command Navigation and User Experience", 3)
    add_image_contain(slide, "frontend/predictive-sidebar-desktop.png", Inches(7.1), Inches(1.25), Inches(5.65), Inches(4.95))
    add_bullets(slide, [
        "Sidebar sections cover Overview, GIS Operations, Predictive Risk, Upstream Watch, River Monitoring, Alerts, Incidents, Citizen Reports, Activity, Simulation, and Rainfall Heatmap.",
        "The top bar exposes operational location selection and current update status.",
        "Responsive screenshots exist for desktop and mobile views, and validation scripts check overflow and rendering.",
        "External Google Flood Hub access is available directly from navigation for official reference mapping.",
    ], Inches(0.72), Inches(1.45), Inches(5.8), Inches(3.0), size=13)
    add_card(slide, Inches(0.72), Inches(4.95), Inches(2.65), Inches(0.92), "Design intent", "Dense command surface for repeated operational use.", BLUE)
    add_card(slide, Inches(3.62), Inches(4.95), Inches(2.65), Inches(0.92), "Operational state", "Live mode, simulation mode, backend errors, and refresh status stay visible.", ORANGE)

    # 4. Overview
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Operations", "Operational Overview Dashboard", 4)
    add_image_contain(slide, "frontend/monitoring-panel-desktop.png", Inches(6.55), Inches(1.22), Inches(6.2), Inches(4.98))
    add_bullets(slide, [
        "National map canvas opens with Pakistan boundary context and overlays command KPIs.",
        "KPI deck summarizes risk, active incidents, reports, resources, monitoring interval, and command statistics.",
        "Hourly analytics ticker provides quick scanning across weather, river, risk, reports, and operational feeds.",
        "Mode banners warn when the system is in simulation or when a scenario is configured.",
    ], Inches(0.72), Inches(1.44), Inches(5.45), Inches(2.8), size=13)
    add_card(slide, Inches(0.72), Inches(4.75), Inches(2.6), Inches(1.0), "Live refresh", "Dashboard batches weather, river, risk, GIS, event, report, alert, incident, and resource requests.", TEAL)
    add_card(slide, Inches(3.58), Inches(4.75), Inches(2.6), Inches(1.0), "Flood cadence", "River layer refreshes every 15 minutes normally and every 3 minutes during flood-like states.", ORANGE)

    # 5. GIS
    slide = prs.slides.add_slide(blank)
    add_header(slide, "GIS Operations", "Layered Map Intelligence", 5)
    add_image_contain(slide, "frontend/river-marker-click.png", Inches(6.35), Inches(1.2), Inches(6.38), Inches(5.05))
    left_items = [
        "Static layers: Pakistan boundary, district names, catchments, glaciers, rivers, lakes, dams, and risk locations.",
        "Dynamic layers: district screening, river status, model forecast points, monitoring gauges, Flood Hub, reports, incidents, and response resources.",
        "Coordinate tools support lat/long entry, map picking, feature lookup, and study mode for corresponding catchment layers.",
        "Kashmir boundary source selector, LoC reference, and explicit display-geometry disclaimers avoid false boundary certainty.",
    ]
    add_bullets(slide, left_items, Inches(0.72), Inches(1.42), Inches(5.3), Inches(3.45), size=12.5)
    add_card(slide, Inches(0.72), Inches(5.2), Inches(2.52), Inches(0.88), "Map engine", "MapLibre with compressed, cached GeoJSON assets.", BLUE)
    add_card(slide, Inches(3.48), Inches(5.2), Inches(2.52), Inches(0.88), "Study mode", "Clips layers to selected HydroBASINS catchment.", TEAL)

    # 6. Rainfall
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Rainfall", "Heatmap and Weather Context", 6)
    add_image_contain(slide, "frontend/chart-1440.png", Inches(6.9), Inches(1.42), Inches(5.45), Inches(2.22))
    add_image_contain(slide, "frontend/upstream-cache-1440.png", Inches(6.9), Inches(4.15), Inches(5.45), Inches(0.75))
    add_bullets(slide, [
        "Rainfall heatmap supports last 1h, 3h, 6h, 24h and forecast 24h, 72h windows.",
        "Open-Meteo weather and regional rainfall samples feed the raster; local data is used when regional refresh is unavailable.",
        "Legend displays rainfall accumulation ranges from light precipitation through 100+ mm.",
        "Rainfall layer can be enabled directly from the rainfall view or as a GIS dynamic layer.",
    ], Inches(0.72), Inches(1.45), Inches(5.8), Inches(3.1), size=13)
    add_card(slide, Inches(0.72), Inches(4.95), Inches(2.75), Inches(0.95), "Source hygiene", "Cached regional rainfall can be shown when provider refresh fails.", ORANGE)
    add_card(slide, Inches(3.72), Inches(4.95), Inches(2.75), Inches(0.95), "Operational use", "Pairs precipitation signal with district risk, river status, and reports.", TEAL)

    # 7. Predictive
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Predictive Risk", "Multi-Hazard Forecast Views", 7)
    add_image_contain(slide, "frontend/predictive-sidebar-desktop.png", Inches(6.85), Inches(1.16), Inches(5.65), Inches(5.15))
    add_bullets(slide, [
        "Riverine short-term: station discharge and flood class for Trimmu and Balloki horizons.",
        "Flash flood: 0-6 hour probability for hill torrents, streams, and urban catchments.",
        "Urban flooding: city grid or ward susceptibility using rainfall, DEM, slope, drainage, land use, and reports.",
        "Seasonal basin risk and upstream or transboundary anomaly detection give wider decision context.",
        "Stored forecast points are charted by model case and selectable target.",
    ], Inches(0.72), Inches(1.42), Inches(5.75), Inches(3.75), size=12.5)
    add_card(slide, Inches(0.72), Inches(5.46), Inches(2.7), Inches(0.82), "Model stance", "Rule-based live baseline remains primary until validated ML artifacts meet deployment gates.", RED)
    add_card(slide, Inches(3.66), Inches(5.46), Inches(2.7), Inches(0.82), "Model families", "XGBoost, LightGBM, TFT comparison, and anomaly methods are represented.", BLUE)

    # 8. Upstream
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Upstream Watch", "Source Flood Forecasting", 8)
    add_image_contain(slide, "frontend/source-tab-expanded.png", Inches(6.26), Inches(1.18), Inches(6.55), Inches(5.15))
    add_bullets(slide, [
        "Station selector covers upstream source stations with river and basin metadata.",
        "GloFAS and GEOGloWS forecasts are shown in parallel with current flow, model peaks, peak timing, and provider notes.",
        "Trend strips summarize retrospective modeled history; rainfall context shows forecast totals and probabilities.",
        "Stored cache timestamp, next refresh time, automatic hourly update, and manual refresh button are exposed to users.",
    ], Inches(0.72), Inches(1.42), Inches(5.08), Inches(3.3), size=12.7)
    add_card(slide, Inches(0.72), Inches(5.15), Inches(2.4), Inches(0.88), "Lead time", "Hours to days for upstream and transboundary signals.", TEAL)
    add_card(slide, Inches(3.34), Inches(5.15), Inches(2.4), Inches(0.88), "Traceability", "Forecast panel labels model source and stale or unavailable data.", ORANGE)

    # 9. River monitoring
    slide = prs.slides.add_slide(blank)
    add_header(slide, "River Monitoring", "Stations, Gauges, and Flow Status", 9)
    add_image_contain(slide, "frontend/river-monitoring-desktop.png", Inches(6.25), Inches(1.18), Inches(6.45), Inches(5.12))
    add_bullets(slide, [
        "Detailed table filters by river, unit, station search, and selected station inspector.",
        "Shows current flow, forecast flow, flood category, 24h change, expected peak, and timestamps.",
        "Separates FFD observed status from GloFAS modeled values and labels measured versus modeled fields.",
        "Map river colors change based on current river status and refresh at operational cadence.",
    ], Inches(0.72), Inches(1.44), Inches(5.15), Inches(3.25), size=12.8)
    add_card(slide, Inches(0.72), Inches(5.1), Inches(2.45), Inches(0.95), "Units", "m3/s and cusecs.", BLUE)
    add_card(slide, Inches(3.38), Inches(5.1), Inches(2.45), Inches(0.95), "Quality", "Observed water height is left unavailable if not published.", RED)

    # 10. Station detail
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Station Detail", "Hydrograph, Hover Cards, and Forecast Windows", 10)
    add_image_contain(slide, "frontend/monitoring-detail-desktop.png", Inches(7.16), Inches(1.12), Inches(5.25), Inches(5.3))
    add_bullets(slide, [
        "Hover cards show compact station facts without triggering heavy fetches on every mouse move.",
        "Pinned station details retain original readings, hydrograph, and upstream network context.",
        "Discharge forecast tab adds a larger chart with observed history, modeled history, forecast days, selectable legend, table, and CSV download.",
        "Observed-history calls are abortable and cached; stale readings and missing thresholds are labeled rather than inferred.",
    ], Inches(0.72), Inches(1.42), Inches(5.95), Inches(3.55), size=12.6)
    add_card(slide, Inches(0.72), Inches(5.35), Inches(2.8), Inches(0.82), "Accessibility", "Keyboard focus proxy, Enter to pin, Escape to dismiss.", TEAL)
    add_card(slide, Inches(3.78), Inches(5.35), Inches(2.8), Inches(0.82), "Scale", "Large networks use native MapLibre layers instead of DOM markers.", BLUE)

    # 11. Alerts and incidents
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Response Operations", "Alerts and Incident Management", 11)
    add_rect(slide, Inches(6.82), Inches(1.22), Inches(5.55), Inches(4.85), fill=PALE, line=LINE, radius=True)
    add_textbox(slide, Inches(7.1), Inches(1.52), Inches(4.95), Inches(0.45),
                "Alert -> Incident -> Activity Trail", size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    steps = [
        ("1", "Threshold or risk signal", BLUE),
        ("2", "Alert record", ORANGE),
        ("3", "Linked incident", RED),
        ("4", "Timeline event", TEAL),
    ]
    for idx, (num, label, color) in enumerate(steps):
        x = Inches(7.28) + Inches(1.18) * idx
        circ = slide.shapes.add_shape(9, x, Inches(2.42), Inches(0.6), Inches(0.6))
        circ.fill.solid()
        circ.fill.fore_color.rgb = color
        circ.line.color.rgb = color
        add_textbox(slide, x, Inches(2.56), Inches(0.6), Inches(0.24), num, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, x - Inches(0.28), Inches(3.22), Inches(1.16), Inches(0.55), label, size=9.5, color=INK, align=PP_ALIGN.CENTER)
    add_bullets(slide, [
        "Alerts table filters all, operational, and simulated-flow alerts.",
        "Each alert records station or location, severity, message, source, timestamp, threshold metadata, and linked incident.",
        "Incident table tracks severity, status, confidence, source, creation time, and detail drawer.",
        "Activity timeline records risk updates, monitoring mode changes, incident creation, citizen reports, confirmations, and required alerts.",
    ], Inches(0.72), Inches(1.45), Inches(5.72), Inches(3.35), size=12.8)
    add_card(slide, Inches(0.72), Inches(5.12), Inches(2.8), Inches(0.88), "Auditability", "Operational and simulated records are visually separated.", ORANGE)
    add_card(slide, Inches(3.78), Inches(5.12), Inches(2.8), Inches(0.88), "Workflow base", "Current screens are ready for approval and escalation layers.", TEAL)

    # 12. Citizen reports
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Field Input", "Citizen Reports and Verification Queue", 12)
    add_image_contain(slide, "frontend/monitoring-tab-desktop.png", Inches(6.75), Inches(1.18), Inches(5.8), Inches(5.15))
    add_bullets(slide, [
        "Reports are fetched from the backend database and newly created reports are saved back through the API.",
        "Full report entry supports operational location, latitude, longitude, Google Maps link parsing, browser location, water depth, source, image URL, and video URL.",
        "Geotagged image upload reads EXIF GPS when present; otherwise image GPS fields remain visible as null.",
        "Reports stay highlighted as live for one minute, appear in a scrollable list, can be marked important, and can be closed in the database.",
    ], Inches(0.72), Inches(1.42), Inches(5.72), Inches(3.58), size=12.5)
    add_card(slide, Inches(0.72), Inches(5.35), Inches(2.75), Inches(0.82), "Validation", "Required description and numeric coordinate/depth checks.", BLUE)
    add_card(slide, Inches(3.72), Inches(5.35), Inches(2.75), Inches(0.82), "Map use", "Google Maps search and coordinates help field localization.", TEAL)

    # 13. Population and exposure
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Exposure", "Flood-Affected Population and Impact Analysis", 13)
    add_bullets(slide, [
        "GIS Operations includes a population panel that loads latest analysis or starts a background calculation job.",
        "Sources include historical satellite water reference, user-provided flood GeoJSON, configured segmentation output, or explicit synthetic sample.",
        "Cards summarize exposed people, percentage of known population, mapped-water coverage with population data, and affected settlements.",
        "Map layers include flood extent, population density, administrative exposure, and affected settlements.",
        "Exposure API also supports roads, buildings, hospitals, schools, settlements, summary artifacts, audits, and layer downloads.",
    ], Inches(0.72), Inches(1.42), Inches(5.88), Inches(4.1), size=12.4)
    add_rect(slide, Inches(6.95), Inches(1.35), Inches(5.45), Inches(4.85), fill=PALE, line=LINE, radius=True)
    add_textbox(slide, Inches(7.25), Inches(1.72), Inches(4.8), Inches(0.36),
                "Trimmu / Jhang Pilot Reference", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_metric(slide, Inches(7.35), Inches(2.42), Inches(2.05), "Exposed people", "313,436", RED)
    add_metric(slide, Inches(9.65), Inches(2.42), Inches(2.05), "Known pop.", "9.66%", ORANGE)
    add_metric(slide, Inches(7.35), Inches(3.56), Inches(2.05), "Water coverage", "57.44%", BLUE)
    add_metric(slide, Inches(9.65), Inches(3.56), Inches(2.05), "Settlements", "16", GREEN)
    add_textbox(slide, Inches(7.35), Inches(4.82), Inches(4.35), Inches(0.52),
                "The deck keeps these values labeled as preliminary satellite reference analysis, not a live forecast or full affected-person count.",
                size=10, color=MUTED, align=PP_ALIGN.CENTER)

    # 14. External intelligence
    slide = prs.slides.add_slide(blank)
    add_header(slide, "External Intelligence", "Google Flood Hub, Water Layers, and Resources", 14)
    add_image_contain(slide, "frontend/google-flood-desktop.png", Inches(6.35), Inches(1.18), Inches(6.35), Inches(5.12))
    add_bullets(slide, [
        "Google Flood Hub side navigation opens the official Flood Hub map in a new browser tab.",
        "GIS toggle displays official Flood Forecasting API river gauge points when credentials are configured.",
        "Markers use severity colors for extreme, danger, warning, normal, and no-data statuses with issue and forecast interval details.",
        "Response-resource layers can plot shelters, hospitals, and emergency contacts from authority-supplied data.",
        "Water intelligence separates GIS water features, station metadata, modeled forecast points, and observed gauge status.",
    ], Inches(0.72), Inches(1.42), Inches(5.3), Inches(3.75), size=12.4)
    add_card(slide, Inches(0.72), Inches(5.45), Inches(2.65), Inches(0.82), "No substitution", "Without Google credentials, no fake forecasts are shown.", RED)
    add_card(slide, Inches(3.62), Inches(5.45), Inches(2.65), Inches(0.82), "Attribution", "Flood Forecasting API data is labeled with source links.", BLUE)

    # 15. Simulation
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Simulation", "Monitoring Scenario Testing", 15)
    add_bullets(slide, [
        "Simulation page loads monitoring stations and lets users select station, measurement type, unit, start time, and gauge reference.",
        "Threshold source can be saved database thresholds or custom test thresholds.",
        "Testing readings support offsets, values, p25/p75 bands, expected level labels, profiles, and up to 240 rows.",
        "Play saves readings, creates alerts and incidents on threshold crossing, and pushes users to the generated records.",
        "Simulation mode is clearly stamped and separated from live operational data.",
    ], Inches(0.72), Inches(1.42), Inches(5.82), Inches(3.8), size=12.6)
    add_rect(slide, Inches(7.02), Inches(1.35), Inches(5.15), Inches(4.85), fill=PALE, line=LINE, radius=True)
    add_textbox(slide, Inches(7.36), Inches(1.74), Inches(4.45), Inches(0.36),
                "Scenario Flow", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    flow = [("Configure", BLUE), ("Add readings", TEAL), ("Play", ORANGE), ("Review alerts", RED)]
    for idx, (label, color) in enumerate(flow):
        y = Inches(2.45) + Inches(0.8) * idx
        add_card(slide, Inches(7.65), y, Inches(3.9), Inches(0.52), label, " ", color)
        arrow = slide.shapes.add_shape(13, Inches(9.37), y + Inches(0.58), Inches(0.28), Inches(0.28))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(168, 181, 194)
        arrow.line.color.rgb = RGBColor(168, 181, 194)
        if idx == len(flow) - 1:
            arrow.fill.transparency = 100
            arrow.line.transparency = 100

    # 16. Platform foundation
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Platform", "Backend, Data, and Quality Foundation", 16)
    add_bullets(slide, [
        "FastAPI backend exposes health, system status, weather, river, flood risk, GIS, exposure, population, reports, incidents, alerts, resources, automation, and simulation routes.",
        "SQLite is default for local development; PostgreSQL is supported through DATABASE_URL.",
        "Live providers include Open-Meteo weather and GloFAS modeled discharge; Google Flood API is optional and credential-gated.",
        "ML artifacts are versioned with model cards; underperforming research models are explicitly not deployed.",
        "Tests cover API behavior, GIS consistency, rainfall, river monitoring, model baselines, exposure, population, station history, simulation, and UI validations.",
    ], Inches(0.72), Inches(1.42), Inches(6.1), Inches(4.05), size=12.3)
    add_rect(slide, Inches(7.18), Inches(1.45), Inches(4.95), Inches(4.45), fill=PALE, line=LINE, radius=True)
    stack = [
        ("React + TypeScript + MapLibre", "Authority dashboard and operational map", BLUE),
        ("FastAPI + SQLAlchemy", "APIs, persistence, simulations, events", TEAL),
        ("Providers + cached assets", "Weather, rivers, rainfall, Google, compressed GeoJSON", ORANGE),
        ("Risk + ML services", "Rule-based live scoring and gated model artifacts", RED),
        ("Tests + scripts", "Backend, GIS, UI, exposure, and population validation", GREEN),
    ]
    for i, (title, body, color) in enumerate(stack):
        add_card(slide, Inches(7.48), Inches(1.78) + Inches(0.76) * i, Inches(4.35), Inches(0.57), title, body, color)

    # 17. Future scope
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Future Scope", "High-Value Enhancements", 17)
    scope_cards = [
        ("Governance workflow", "Role-based login, report verification queues, approval chain, escalation to higher authorities, and signed action logs.", BLUE),
        ("Warning dissemination", "SMS, WhatsApp, email, siren, CAP/alerting feeds, multilingual templates, and delivery confirmation.", RED),
        ("Advanced modeling", "Validated ensemble models, TFT sequence forecasts, hydraulic depth via HEC-RAS 2D, uncertainty bands, and model drift monitoring.", TEAL),
        ("IoT and field feeds", "Gauge telemetry, low-cost rain stations, reservoir releases, rescue-team mobile updates, and offline-first data capture.", ORANGE),
        ("Damage and exposure", "Near-real-time flood segmentation, infrastructure impact, evacuation routing, shelter capacity, and vulnerable-population overlays.", GREEN),
        ("Command operations", "Tasking, resource dispatch, logistics stock, mutual-aid requests, post-event reports, and KPI dashboards for leadership.", BLUE),
    ]
    for i, (title, body, color) in enumerate(scope_cards):
        x = Inches(0.72) + Inches(4.12) * (i % 3)
        y = Inches(1.45) + Inches(1.72) * (i // 3)
        add_card(slide, x, y, Inches(3.72), Inches(1.34), title, body, color)
    add_textbox(slide, Inches(0.82), Inches(5.35), Inches(11.7), Inches(0.5),
                "Priority principle: strengthen trust, provenance, and operational accountability before expanding automated public warning.",
                size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    # 18. Roadmap
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Roadmap", "Suggested Implementation Path", 18)
    add_timeline(slide, [
        ("0-3 months", "Harden current portal", [
            "Auth and user roles",
            "Report approval workflow",
            "Alert templates and audit trail",
            "Production deployment checklist",
        ], BLUE),
        ("3-6 months", "Operational integrations", [
            "Authority resource database",
            "Notification channels",
            "Telemetry connectors",
            "Incident tasking and dispatch",
        ], TEAL),
        ("6-12 months", "Model and exposure expansion", [
            "Validated ML deployment gate",
            "Hydraulic depth studies",
            "More regions for exposure",
            "Automated quality dashboards",
        ], ORANGE),
        ("12+ months", "National command maturity", [
            "Multi-agency operations room",
            "Mobile/offline field apps",
            "Public warning publishing",
            "After-action analytics",
        ], RED),
    ], Inches(0.72), Inches(1.55), Inches(11.9))
    add_textbox(slide, Inches(0.8), Inches(5.45), Inches(11.7), Inches(0.55),
                "The portal already has a strong technical foundation. The next leap is workflow maturity: verified data, controlled approvals, reliable integrations, and trusted alerts.",
                size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
